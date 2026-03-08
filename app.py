from __future__ import annotations

import io
import os
from dataclasses import dataclass

from flask import Flask, render_template, request, send_file
from openai import OpenAI
from pypdf import PdfReader
from werkzeug.datastructures import FileStorage
from docx import Document

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024  # 25 MB


@dataclass
class ProcessResult:
    source_text: str = ""
    translated_text: str = ""
    error: str = ""


def extract_text_from_file(upload: FileStorage) -> str:
    filename = (upload.filename or "").lower()

    if filename.endswith(".txt"):
        return upload.read().decode("utf-8", errors="ignore")

    if filename.endswith(".docx"):
        doc = Document(upload)
        return "\n".join(p.text for p in doc.paragraphs)

    if filename.endswith(".pdf"):
        reader = PdfReader(upload)
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages)

    raise ValueError("Định dạng file chưa được hỗ trợ. Vui lòng dùng .txt, .docx hoặc .pdf")


def translate_and_polish(text: str) -> str:
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("Thiếu OPENAI_API_KEY. Hãy cấu hình biến môi trường trước khi dịch.")

    client = OpenAI(api_key=api_key)
    response = client.responses.create(
        model=os.getenv("OPENAI_MODEL", "gpt-4.1-mini"),
        input=[
            {
                "role": "system",
                "content": (
                    "Bạn là biên dịch viên tiếng Anh sang tiếng Việt chuyên nghiệp. "
                    "Nhiệm vụ: dịch chính xác toàn bộ nội dung đầu vào sang tiếng Việt, "
                    "sau đó biên tập văn phong tự nhiên, mượt mà, dễ hiểu. "
                    "Không thêm thông tin ngoài nội dung gốc. "
                    "Chỉ trả về văn bản tiếng Việt cuối cùng."
                ),
            },
            {"role": "user", "content": text},
        ],
    )
    return response.output_text.strip()


@app.route("/", methods=["GET"])
def index() -> str:
    return render_template("index.html", result=ProcessResult())


@app.route("/process", methods=["POST"])
def process() -> str:
    upload = request.files.get("file")
    if not upload or not upload.filename:
        return render_template(
            "index.html",
            result=ProcessResult(error="Vui lòng tải lên một file .txt, .docx hoặc .pdf"),
        )

    try:
        source_text = extract_text_from_file(upload)
        if not source_text.strip():
            raise ValueError("Không đọc được nội dung văn bản từ file đã tải lên.")

        translated_text = translate_and_polish(source_text)
        result = ProcessResult(source_text=source_text, translated_text=translated_text)
        return render_template("index.html", result=result)
    except Exception as exc:  # noqa: BLE001
        return render_template("index.html", result=ProcessResult(error=str(exc)))


@app.route("/download", methods=["POST"])
def download():
    edited_text = request.form.get("edited_text", "")
    output_format = request.form.get("output_format", "txt")

    if output_format == "docx":
        doc = Document()
        for paragraph in edited_text.split("\n"):
            doc.add_paragraph(paragraph)

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return send_file(
            buffer,
            as_attachment=True,
            download_name="ban-dich-tieng-viet.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    buffer = io.BytesIO(edited_text.encode("utf-8"))
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name="ban-dich-tieng-viet.txt",
        mimetype="text/plain; charset=utf-8",
    )


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8000, debug=True)
